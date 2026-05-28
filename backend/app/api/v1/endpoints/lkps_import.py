import uuid
from datetime import date, datetime
from io import BytesIO

import openpyxl
from fastapi import APIRouter, Depends, File, HTTPException, UploadFile, status
from sqlalchemy.orm import Session

from app.db import get_db
from app.models.lkps import (
    LkpsBebanKerjaDosen,
    LkpsCapstoneDesign,
    LkpsDosenProfil,
    LkpsIntegrasiPenelitian,
    LkpsIpkLulusan,
    LkpsK3lDokumen,
    LkpsK3lFasilitas,
    LkpsKerjasama,
    LkpsKesesuaianKerja,
    LkpsKinerjaDtps,
    LkpsKurikulum,
    LkpsMahasiswaAktif,
    LkpsMasaStudi,
    LkpsMataKuliahPpi,
    LkpsMkBasicScience,
    LkpsPembimbingLapangan,
    LkpsPenggunaanDana,
    LkpsPenelitianMahasiswa,
    LkpsPenelitianSummary,
    LkpsPkmSummary,
    LkpsPppiDisiplin,
    LkpsPrasarana,
    LkpsPrestasiMahasiswa,
    LkpsProdukJasa,
    LkpsPublikasiIlmiah,
    LkpsRekognisiDtps,
    LkpsSitasiDtps,
    LkpsSpmiDokumen,
    LkpsSpmiPelaksanaan,
    LkpsSubmission,
    LkpsTempatKerja,
    LkpsTenagaKependidikan,
    LkpsVmts,
    LkpsWaktuTunggu,
    LkpsKepuasanPengguna,
    LkpsLuaranPenelitian,
)
from app.models.user import User
from app.utils.dependencies import get_current_user

router = APIRouter(prefix="/lkps/import", tags=["lkps-import"])

# ---------------------------------------------------------------------------
# Scalar conversion helpers
# ---------------------------------------------------------------------------

def _s(v) -> str | None:
    """Coerce to stripped string or None."""
    return str(v).strip() if v is not None else None


def _i(v) -> int | None:
    """Coerce to int or None."""
    if v is None:
        return None
    try:
        return int(v)
    except (TypeError, ValueError):
        return None


def _f(v) -> float | None:
    """Coerce to float or None."""
    if v is None:
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def _d(v) -> date | None:
    """Parse a date value from openpyxl (datetime, date, or string)."""
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    # Try parsing common string formats
    if isinstance(v, str):
        v = v.strip()
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
            try:
                return datetime.strptime(v, fmt).date()
            except ValueError:
                continue
    return None


def _b(v) -> bool:
    """Return True if the cell value is the checkmark 'V' (case-insensitive)."""
    if v is None:
        return False
    return str(v).strip().upper() == "V"


def _tingkat(v_intr, v_nas, v_lokal) -> str | None:
    """Return the tingkat string for whichever column holds 'V'."""
    if _b(v_intr):
        return "internasional"
    if _b(v_nas):
        return "nasional"
    if _b(v_lokal):
        return "lokal"
    return None


def _col(row_values: tuple, letter: str):
    """Return the zero-indexed value from a row_values tuple by Excel column letter."""
    idx = 0
    for ch in letter.upper():
        idx = idx * 26 + (ord(ch) - ord("A") + 1)
    return row_values[idx - 1]


def _has_data(*values) -> bool:
    """Return True when at least one imported input cell is intentionally filled."""
    for value in values:
        if value is None:
            continue
        if isinstance(value, str) and value.strip() == "":
            continue
        return True
    return False


def _carry_forward(value, previous):
    """Use the current cell value, or the previous non-empty category value."""
    current = _s(value)
    return current if current else previous


# ---------------------------------------------------------------------------
# Fixed-row key maps (mirrors the export mappers exactly)
# ---------------------------------------------------------------------------

_DANA_ROW: dict[str, int] = {
    "biaya_dosen":               7,
    "biaya_tendik":              8,
    "biaya_op_pembelajaran":     9,
    "biaya_op_tidak_langsung":  10,
    "biaya_praktik_ppi":        11,
    "biaya_investasi":          12,
    "biaya_kemahasiswaan":      13,
    "biaya_penelitian":         15,
    "biaya_pkm":                16,
}
_DANA_ROW_INV: dict[int, str] = {v: k for k, v in _DANA_ROW.items()}

_PENELITIAN_ROW: dict[str, int] = {
    "perguruan_tinggi_mandiri": 9,
    "dalam_negeri":            10,
    "luar_negeri":             11,
}
_PENELITIAN_ROW_INV: dict[int, str] = {v: k for k, v in _PENELITIAN_ROW.items()}

_PKM_ROW: dict[str, int] = {
    "perguruan_tinggi_mandiri": 10,
    "dalam_negeri":            11,
    "luar_negeri":             12,
}
_PKM_ROW_INV: dict[int, str] = {v: k for k, v in _PKM_ROW.items()}

_PUBLIKASI_ROW_AKADEMIK: dict[int, str] = {
    7:  "jurnal_nasional_tidak_terakreditasi",
    8:  "jurnal_nasional_terakreditasi",
    9:  "jurnal_internasional",
    10: "jurnal_internasional_bereputasi",
    11: "prosiding_nasional",
    12: "prosiding_internasional_tidak_terindeks",
    13: "prosiding_internasional_terindeks",
}

_PUBLIKASI_ROW_VOKASI: dict[int, str] = {
    **_PUBLIKASI_ROW_AKADEMIK,
    14: "pagelaran_wilayah",
    15: "pagelaran_nasional",
    16: "pagelaran_internasional",
}

_LUARAN_START: dict[str, int] = {
    "4f-1":   7,  "4f-2":   7,  "4f-3":  16,  "4f-4":   7,
    "6e3-1": 12, "6e3-2":   8, "6e3-3":  16, "6e3-4":   8,
}

_LUARAN_SHEET_META: dict[str, tuple[str, str]] = {
    # sheet_name: (sumber, jenis_luaran)
    "4f-1":   ("dtps",      "paten"),
    "4f-2":   ("dtps",      "hak_cipta"),
    "4f-3":   ("dtps",      "teknologi"),
    "4f-4":   ("dtps",      "buku"),
    "6e3-1":  ("mahasiswa", "paten"),
    "6e3-2":  ("mahasiswa", "hak_cipta"),
    "6e3-3":  ("mahasiswa", "teknologi"),
    "6e3-4":  ("mahasiswa", "buku"),
}

_PENDIDIKAN_COL: dict[str, str] = {
    "S3": "C", "S2": "D", "S1": "E",
    "D4": "F", "D3": "G", "D2": "H", "D1": "I", "SMA_SMK": "J",
}
_PENDIDIKAN_COL_INV: dict[str, str] = {v: k for k, v in _PENDIDIKAN_COL.items()}

_IPK_ROW_INV: dict[int, str] = {6: "TS-2", 7: "TS-1", 8: "TS"}

_WT_SECTION_START: dict[str, int] = {
    "D1":  7, "D2": 15, "D3": 23,
    "S1": 31, "S1Tr": 39, "PPI": 47,
}
_WT_TAHUN_OFFSET_INV: dict[int, str] = {0: "TS-2", 1: "TS-1"}

_LULUSAN_ROW_INV: dict[int, str] = {7: "TS-2", 8: "TS-1"}

_KEPUASAN_ROW_INV: dict[int, str] = {
    7:  "Etika",
    8:  "Keahlian pada bidang ilmu (kompetensi utama)",
    9:  "Kemampuan berbahasa asing",
    10: "Penggunaan teknologi informasi",
    11: "Kemampuan berkomunikasi",
    12: "Kerjasama tim",
    13: "Pengembangan diri",
}

_SPMI_DOK_ROW_INV: dict[int, str] = {
    5: "Kebijakan SPMI",
    6: "Pedoman penerapan siklus PPEPP standar pendidikan tinggi dalam SPMI",
    7: "Standar dan/atau kriteria, norma, acuan mutu penyelenggaraan pendidikan dan pengelolaan perguruan tinggi",
    8: "Tata cara pendokumentasian implementasi SPMI",
}

_PPEPP_ROW_INV: dict[int, str] = {
    5: "Penetapan",
    6: "Pelaksanaan",
    7: "Evaluasi",
    8: "Pengendalian",
    9: "Peningkatan",
}

_MASA_STUDI_SECTIONS: dict[str, tuple[int, list[str]]] = {
    "D1":   (7,  ["TS-1", "TS"]),
    "D2":   (14, ["TS-3", "TS-2", "TS-1", "TS"]),
    "D3":   (23, ["TS-5", "TS-4", "TS-3", "TS-2", "TS-1", "TS"]),
    "S1":   (34, ["TS-7", "TS-6", "TS-5", "TS-4", "TS-3", "TS-2", "TS-1", "TS"]),
    "S1Tr": (34, ["TS-7", "TS-6", "TS-5", "TS-4", "TS-3", "TS-2", "TS-1", "TS"]),
    "S2":   (47, ["TS-3", "TS-2", "TS-1", "TS"]),
    "S2Tr": (47, ["TS-3", "TS-2", "TS-1", "TS"]),
    "S3":   (56, ["TS-5", "TS-4", "TS-3", "TS-2", "TS-1", "TS"]),
    "S3Tr": (56, ["TS-5", "TS-4", "TS-3", "TS-2", "TS-1", "TS"]),
    "PPI":  (67, ["TS-2", "TS-1", "TS"]),
}

_PPPI_DISIPLIN_ROW_INV: dict[int, str] = {
    17: "Kebumian dan Energi",
    18: "Rekayasa Sipil dan Lingkungan Terbangun",
    19: "Industri",
    20: "Konservasi dan Pengelolaan Sumber Daya Alam",
    21: "Pertanian dan Hasil Pertanian",
    22: "Teknologi Kelautan dan Perkapalan",
    23: "Aeronotika dan Astronotika",
}


# ---------------------------------------------------------------------------
# Main endpoint
# ---------------------------------------------------------------------------

@router.post(
    "/{submission_id}",
    status_code=status.HTTP_200_OK,
    summary="Import LKPS data from Excel",
)
async def import_lkps_excel(
    submission_id: uuid.UUID,
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """
    Mengimpor data input LKPS dari file Excel (T-49).
    Membaca file .xlsx yang diunggah dan memasukkan datanya ke dalam database
    berdasarkan submission_id yang dipilih.
    """
    if not file.filename.endswith((".xlsx", ".xls")):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Format file tidak valid. Harap unggah file Excel (.xlsx).",
        )

    # Validate submission exists
    submission = (
        db.query(LkpsSubmission).filter(LkpsSubmission.id == submission_id).first()
    )
    if not submission:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="LKPS Submission tidak ditemukan.",
        )

    # Read and load the workbook
    try:
        contents = await file.read()
        wb = openpyxl.load_workbook(filename=BytesIO(contents), data_only=True)
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Gagal membaca file Excel: {str(e)}",
        )

    try:
        _import_all_sheets(wb, db, submission_id)
        db.commit()
    except HTTPException:
        db.rollback()
        raise
    except Exception as e:
        db.rollback()
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Gagal menyimpan data ke database: {str(e)}",
        )

    return {
        "message": "Data LKPS berhasil diimpor dari Excel.",
        "submission_id": submission_id,
    }


# ---------------------------------------------------------------------------
# Dispatcher
# ---------------------------------------------------------------------------

def _import_all_sheets(wb, db: Session, sid: uuid.UUID) -> None:
    _import_sheet_pppi(wb, db, sid)
    _import_sheet_1(wb, db, sid)
    _import_sheet_2a(wb, db, sid)
    _import_sheet_2b(wb, db, sid)
    _import_sheet_3a1(wb, db, sid)
    _import_sheet_3a2(wb, db, sid)
    _import_sheet_3a3(wb, db, sid)
    _import_sheet_3a4(wb, db, sid)
    _import_sheet_3a5(wb, db, sid)
    _import_sheet_3b(wb, db, sid)
    _import_sheet_3c(wb, db, sid)
    _import_sheet_4a(wb, db, sid)
    _import_sheet_4b(wb, db, sid)
    _import_sheet_4c(wb, db, sid)
    _import_sheet_4d(wb, db, sid)
    _import_sheet_4e(wb, db, sid)
    _import_luaran_sheets(wb, db, sid)
    _import_sheet_4g(wb, db, sid)
    _import_sheet_4h(wb, db, sid)
    _import_sheet_4i(wb, db, sid)
    _import_sheet_4j(wb, db, sid)
    _import_sheet_4k(wb, db, sid)
    _import_sheet_5a(wb, db, sid)
    _import_sheet_5b(wb, db, sid)
    _import_sheet_5c(wb, db, sid)
    _import_sheet_6a(wb, db, sid)
    _import_sheet_6b(wb, db, sid)
    _import_sheet_6c(wb, db, sid)
    _import_sheet_6d(wb, db, sid)
    _import_sheet_6e1(wb, db, sid)
    _import_sheet_6e2(wb, db, sid)
    _import_sheet_6e4(wb, db, sid)
    _import_sheet_6f1(wb, db, sid)
    _import_sheet_6f2(wb, db, sid)
    _import_sheet_6g1(wb, db, sid)
    _import_sheet_6g2(wb, db, sid)
    _import_sheet_6h1(wb, db, sid)
    _import_sheet_6h2(wb, db, sid)
    _import_sheet_6i(wb, db, sid)
    _import_sheet_7a(wb, db, sid)
    _import_sheet_7b(wb, db, sid)


# ---------------------------------------------------------------------------
# PPI Profile – Disiplin Teknik Keinsinyuran
# ---------------------------------------------------------------------------

def _import_sheet_pppi(wb, db: Session, sid: uuid.UUID) -> None:
    if "PSPPI" not in wb.sheetnames:
        return
    ws = wb["PSPPI"]
    db.query(LkpsPppiDisiplin).filter(LkpsPppiDisiplin.submission_id == sid).delete()
    no_counter = 1
    for row_num, disiplin in _PPPI_DISIPLIN_ROW_INV.items():
        row = ws[row_num]
        # C=Ya, D=Tidak. Only import rows where the user checked one option.
        if not _has_data(row[2].value, row[3].value):
            continue
        db.add(LkpsPppiDisiplin(
            submission_id=sid,
            no=no_counter,
            disiplin=disiplin,
            diselenggarakan=_b(row[2].value),
        ))
        no_counter += 1
    db.flush()


# ---------------------------------------------------------------------------
# Section 1 – VMTS
# ---------------------------------------------------------------------------

def _import_sheet_1(wb, db: Session, sid: uuid.UUID) -> None:
    if "1" not in wb.sheetnames:
        return
    ws = wb["1"]
    db.query(LkpsVmts).filter(LkpsVmts.submission_id == sid).delete()
    no_counter = 1
    current_jenis_vmts = None
    for row in ws.iter_rows(min_row=7, values_only=True):
        # Col A=no, B=jenis_vmts, C=pernyataan, D=no_sk, E=link_dokumen
        if not _has_data(row[2], row[3], row[4]):
            continue
        current_jenis_vmts = _carry_forward(row[1], current_jenis_vmts)
        db.add(LkpsVmts(
            submission_id=sid,
            no=_i(row[0]) or no_counter,
            jenis_vmts=current_jenis_vmts,
            pernyataan=_s(row[2]),
            no_sk=_s(row[3]),
            link_dokumen=_s(row[4]),
        ))
        no_counter += 1
    db.flush()


# ---------------------------------------------------------------------------
# Section 2a – Kerjasama (2a1 pendidikan, 2a2 penelitian, 2a3 pkm)
# ---------------------------------------------------------------------------

_KERJASAMA_DATA_START = {"2a1": 13, "2a2": 12, "2a3": 12}
_KERJASAMA_JENIS = {"2a1": "pendidikan", "2a2": "penelitian", "2a3": "pkm"}


def _import_sheet_2a(wb, db: Session, sid: uuid.UUID) -> None:
    for sheet_name, jenis in _KERJASAMA_JENIS.items():
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        db.query(LkpsKerjasama).filter(
            LkpsKerjasama.submission_id == sid,
            LkpsKerjasama.jenis == jenis,
        ).delete()
        start_row = _KERJASAMA_DATA_START[sheet_name]
        current_lembaga = None
        current_tingkat = None
        for row in ws.iter_rows(min_row=start_row, values_only=True):
            # B=lembaga_mitra, C=internasional, D=nasional, E=lokal,
            # F=judul_kegiatan, G=manfaat, H=tanggal_awal, I=tanggal_akhir,
            # J=durasi(formula), K=status(formula), L=bukti_kerjasama
            if not _has_data(row[1], row[2], row[3], row[4], row[5], row[6], row[7], row[8], row[11]):
                continue
            lembaga = _carry_forward(row[1], current_lembaga)   # col B = index 1
            if not lembaga:
                continue
            current_lembaga = lembaga
            current_tingkat = _tingkat(row[2], row[3], row[4]) or current_tingkat
            db.add(LkpsKerjasama(
                submission_id=sid,
                jenis=jenis,
                lembaga_mitra=lembaga,
                tingkat=current_tingkat,  # C, D, E
                judul_kegiatan=_s(row[5]),
                manfaat=_s(row[6]),
                tanggal_awal=_d(row[7]),
                tanggal_akhir=_d(row[8]),
                bukti_kerjasama=_s(row[11]),  # col L = index 11
            ))
    db.flush()


# ---------------------------------------------------------------------------
# Section 2b – Penggunaan Dana (fixed rows)
# ---------------------------------------------------------------------------

def _import_sheet_2b(wb, db: Session, sid: uuid.UUID) -> None:
    if "2b" not in wb.sheetnames:
        return
    ws = wb["2b"]
    # Delete all existing records for this submission
    db.query(LkpsPenggunaanDana).filter(
        LkpsPenggunaanDana.submission_id == sid
    ).delete()
    for row_num, kode in _DANA_ROW_INV.items():
        row = ws[row_num]
        # C=upps_ts2, D=upps_ts1, E=upps_ts, G=ps_ts2, H=ps_ts1, I=ps_ts
        def cv(cell):
            return cell.value
        if not _has_data(cv(row[2]), cv(row[3]), cv(row[4]), cv(row[6]), cv(row[7]), cv(row[8])):
            continue
        db.add(LkpsPenggunaanDana(
            submission_id=sid,
            kode=kode,
            upps_ts2=_i(cv(row[2])),   # col C = index 2
            upps_ts1=_i(cv(row[3])),   # col D = index 3
            upps_ts=_i(cv(row[4])),    # col E = index 4
            ps_ts2=_i(cv(row[6])),     # col G = index 6
            ps_ts1=_i(cv(row[7])),     # col H = index 7
            ps_ts=_i(cv(row[8])),      # col I = index 8
        ))
    db.flush()


# ---------------------------------------------------------------------------
# Section 3a1 – Kurikulum
# ---------------------------------------------------------------------------

def _import_sheet_3a1(wb, db: Session, sid: uuid.UUID) -> None:
    if "3a1" not in wb.sheetnames:
        return
    ws = wb["3a1"]
    db.query(LkpsKurikulum).filter(LkpsKurikulum.submission_id == sid).delete()
    no_counter = 1
    for row in ws.iter_rows(min_row=10, values_only=True):
        # B=semester, C=kode_mk, D=nama_mk, E=kompetensi, F=sks_kuliah,
        # G=sks_seminar, H=sks_praktikum, I=konversi_jam, J=dokumen_rps, K=unit_penyelenggara
        nama_mk = _s(row[3])  # col D = index 3
        if not nama_mk:
            continue
        db.add(LkpsKurikulum(
            submission_id=sid,
            no=no_counter,
            semester=_i(row[1]),
            kode_mk=_s(row[2]),
            nama_mk=nama_mk,
            kompetensi=_s(row[4]),
            sks_kuliah=_f(row[5]),
            sks_seminar=_f(row[6]),
            sks_praktikum=_f(row[7]),
            konversi_jam=_f(row[8]),
            dokumen_rps=_s(row[9]),
            unit_penyelenggara=_s(row[10]),
        ))
        no_counter += 1
    db.flush()


# ---------------------------------------------------------------------------
# Section 3a2 – Mata Kuliah dan Dokumen Pembelajaran PPI
# ---------------------------------------------------------------------------

def _import_sheet_3a2(wb, db: Session, sid: uuid.UUID) -> None:
    if "3a2" not in wb.sheetnames:
        return
    ws = wb["3a2"]
    db.query(LkpsMataKuliahPpi).filter(LkpsMataKuliahPpi.submission_id == sid).delete()
    no_counter = 1
    for row in ws.iter_rows(min_row=10, values_only=True):
        # B=mata_kuliah, C=bobot_sks, D=konversi_teori_jam,
        # E=konversi_praktik_jam, F=dokumen_rps
        if not _has_data(row[1], row[2], row[3], row[4], row[5]):
            continue
        db.add(LkpsMataKuliahPpi(
            submission_id=sid,
            no=no_counter,
            mata_kuliah=_s(row[1]),
            bobot_sks=_f(row[2]),
            konversi_teori_jam=_f(row[3]),
            konversi_praktik_jam=_f(row[4]),
            dokumen_rps=_s(row[5]),
        ))
        no_counter += 1
    db.flush()


# ---------------------------------------------------------------------------
# Section 3a3 – Integrasi Penelitian/PkM
# ---------------------------------------------------------------------------

def _import_sheet_3a3(wb, db: Session, sid: uuid.UUID) -> None:
    if "3a3" not in wb.sheetnames:
        return
    ws = wb["3a3"]
    db.query(LkpsIntegrasiPenelitian).filter(
        LkpsIntegrasiPenelitian.submission_id == sid
    ).delete()
    for row in ws.iter_rows(min_row=13, values_only=True):
        # B=nama_dosen, C=judul_penelitian_pkm, D=mata_kuliah, E=bentuk_integrasi,
        # F=tahun_ts2(V), G=tahun_ts1(V), H=tahun_ts(V), I=kesesuaian_peta_jalan,
        # J=bukti_sahih, K=kesesuaian_rps
        nama_dosen = _s(row[1])
        if not nama_dosen:
            continue
        db.add(LkpsIntegrasiPenelitian(
            submission_id=sid,
            nama_dosen=nama_dosen,
            judul_penelitian_pkm=_s(row[2]),
            mata_kuliah=_s(row[3]),
            bentuk_integrasi=_s(row[4]),
            tahun_ts2=_b(row[5]),
            tahun_ts1=_b(row[6]),
            tahun_ts=_b(row[7]),
            kesesuaian_peta_jalan=_s(row[8]),
            bukti_sahih=_s(row[9]),
            kesesuaian_rps=_s(row[10]),
        ))
    db.flush()


# ---------------------------------------------------------------------------
# Section 3a4 – MK Basic Science
# ---------------------------------------------------------------------------

def _import_sheet_3a4(wb, db: Session, sid: uuid.UUID) -> None:
    if "3a4" not in wb.sheetnames:
        return
    ws = wb["3a4"]
    db.query(LkpsMkBasicScience).filter(
        LkpsMkBasicScience.submission_id == sid
    ).delete()
    no_counter = 1
    for row in ws.iter_rows(min_row=10, values_only=True):
        # B=nama_mk, C=semester, D=jumlah_sks
        nama_mk = _s(row[1])
        if not nama_mk:
            continue
        db.add(LkpsMkBasicScience(
            submission_id=sid,
            no=no_counter,
            nama_mk=nama_mk,
            semester=_i(row[2]),
            jumlah_sks=_f(row[3]),
        ))
        no_counter += 1
    db.flush()


# ---------------------------------------------------------------------------
# Section 3a5 – Capstone Design
# ---------------------------------------------------------------------------

def _import_sheet_3a5(wb, db: Session, sid: uuid.UUID) -> None:
    if "3a5" not in wb.sheetnames:
        return
    ws = wb["3a5"]
    db.query(LkpsCapstoneDesign).filter(
        LkpsCapstoneDesign.submission_id == sid
    ).delete()
    no_counter = 1
    for row in ws.iter_rows(min_row=10, values_only=True):
        # B=nama_mk_pendukung, C=sks_pendukung, D=nama_mk_capstone,
        # E=sks_capstone, F=semester, G=cakupan_bahasan
        nama_mk_pendukung = _s(row[1])
        if not nama_mk_pendukung:
            continue
        db.add(LkpsCapstoneDesign(
            submission_id=sid,
            no=no_counter,
            nama_mk_pendukung=nama_mk_pendukung,
            sks_pendukung=_f(row[2]),
            nama_mk_capstone=_s(row[3]),
            sks_capstone=_f(row[4]),
            semester=_i(row[5]),
            cakupan_bahasan=_s(row[6]),
        ))
        no_counter += 1
    db.flush()


# ---------------------------------------------------------------------------
# Section 3b – Penelitian Summary (fixed rows)
# ---------------------------------------------------------------------------

def _import_sheet_3b(wb, db: Session, sid: uuid.UUID) -> None:
    if "3b" not in wb.sheetnames:
        return
    ws = wb["3b"]
    db.query(LkpsPenelitianSummary).filter(
        LkpsPenelitianSummary.submission_id == sid
    ).delete()
    for row_num, kode_sumber in _PENELITIAN_ROW_INV.items():
        row = ws[row_num]
        # C=ts2, D=ts1, E=ts
        if not _has_data(row[2].value, row[3].value, row[4].value):
            continue
        db.add(LkpsPenelitianSummary(
            submission_id=sid,
            kode_sumber=kode_sumber,
            ts2=_i(row[2].value),
            ts1=_i(row[3].value),
            ts=_i(row[4].value),
        ))
    db.flush()


# ---------------------------------------------------------------------------
# Section 3c – PkM Summary (fixed rows)
# ---------------------------------------------------------------------------

def _import_sheet_3c(wb, db: Session, sid: uuid.UUID) -> None:
    if "3c" not in wb.sheetnames:
        return
    ws = wb["3c"]
    db.query(LkpsPkmSummary).filter(
        LkpsPkmSummary.submission_id == sid
    ).delete()
    for row_num, kode_sumber in _PKM_ROW_INV.items():
        row = ws[row_num]
        # C=ts2, D=ts1, E=ts
        if not _has_data(row[2].value, row[3].value, row[4].value):
            continue
        db.add(LkpsPkmSummary(
            submission_id=sid,
            kode_sumber=kode_sumber,
            ts2=_i(row[2].value),
            ts1=_i(row[3].value),
            ts=_i(row[4].value),
        ))
    db.flush()


# ---------------------------------------------------------------------------
# Section 4a – Dosen Profil
# ---------------------------------------------------------------------------

def _import_sheet_4a(wb, db: Session, sid: uuid.UUID) -> None:
    if "4a" not in wb.sheetnames:
        return
    ws = wb["4a"]
    db.query(LkpsDosenProfil).filter(LkpsDosenProfil.submission_id == sid).delete()
    no_counter = 1
    current_kategori = None
    for row in ws.iter_rows(min_row=14, values_only=True):
        # B=nama_dosen, C=nidn_nidk, D=kategori, E=prodi_sarjana, F=prodi_magister,
        # G=prodi_doktor, H=bidang_keahlian, I=perusahaan_industri,
        # J=kesesuaian_kompetensi, K=jabatan_akademik, L=no_sertifikat_pendidik,
        # M=bidang_sertifikasi, N=lembaga_penerbit_sertifikasi, O=skip, P=stri,
        # Q=mk_diampu_ps_diakreditasi, R=kesesuaian_bidang_mk, S=mk_diampu_ps_lain
        nama_dosen = _s(row[1])
        if not nama_dosen:
            continue
        current_kategori = _carry_forward(row[3], current_kategori)
        db.add(LkpsDosenProfil(
            submission_id=sid,
            no=no_counter,
            nama_dosen=nama_dosen,
            nidn_nidk=_s(row[2]),
            kategori=current_kategori,
            prodi_sarjana=_s(row[4]),
            prodi_magister=_s(row[5]),
            prodi_doktor=_s(row[6]),
            bidang_keahlian=_s(row[7]),
            perusahaan_industri=_s(row[8]),
            kesesuaian_kompetensi=_s(row[9]),
            jabatan_akademik=_s(row[10]),
            no_sertifikat_pendidik=_s(row[11]),
            bidang_sertifikasi=_s(row[12]),
            lembaga_penerbit_sertifikasi=_s(row[13]),
            skip=_s(row[14]),
            stri=_s(row[15]),
            mk_diampu_ps_diakreditasi=_s(row[16]),
            kesesuaian_bidang_mk=_s(row[17]),
            mk_diampu_ps_lain=_s(row[18]),
        ))
        no_counter += 1
    db.flush()


# ---------------------------------------------------------------------------
# Section 4b – Tenaga Kependidikan
# ---------------------------------------------------------------------------

def _import_sheet_4b(wb, db: Session, sid: uuid.UUID) -> None:
    if "4b" not in wb.sheetnames:
        return
    ws = wb["4b"]
    db.query(LkpsTenagaKependidikan).filter(
        LkpsTenagaKependidikan.submission_id == sid
    ).delete()
    no_counter = 1
    current_pendidikan = None
    current_unit_kerja = None
    for row in ws.iter_rows(min_row=11, values_only=True):
        # B=nama, C=S3, D=S2, E=S1, F=D4, G=D3, H=D2, I=D1, J=SMA_SMK,
        # K=sertifikat_kompetensi, L=unit_kerja
        nama = _s(row[1])
        if not nama:
            continue
        # Find whichever pendidikan column is "V"
        pendidikan = None
        for col_letter, level in _PENDIDIKAN_COL_INV.items():
            col_idx = ord(col_letter) - ord("A")
            if _b(row[col_idx]):
                pendidikan = level
                break
        current_pendidikan = pendidikan or current_pendidikan
        current_unit_kerja = _carry_forward(row[11], current_unit_kerja)
        db.add(LkpsTenagaKependidikan(
            submission_id=sid,
            no=no_counter,
            nama=nama,
            pendidikan_terakhir=current_pendidikan,
            sertifikat_kompetensi=_s(row[10]),  # col K
            unit_kerja=current_unit_kerja,        # col L
        ))
        no_counter += 1
    db.flush()


# ---------------------------------------------------------------------------
# Section 4c – Beban Kerja Dosen
# ---------------------------------------------------------------------------

def _import_sheet_4c(wb, db: Session, sid: uuid.UUID) -> None:
    if "4c" not in wb.sheetnames:
        return
    ws = wb["4c"]
    db.query(LkpsBebanKerjaDosen).filter(
        LkpsBebanKerjaDosen.submission_id == sid
    ).delete()
    no_counter = 1
    for row in ws.iter_rows(min_row=11, values_only=True):
        # B=nama_dosen, C=dtps(V), D=bk_ps_diakreditasi, E=bk_ps_lain_pt,
        # F=bk_ps_luar_pt, G=bk_penelitian, H=bk_pkm, I=bk_tugas_tambahan
        nama_dosen = _s(row[1])
        if not nama_dosen:
            continue
        db.add(LkpsBebanKerjaDosen(
            submission_id=sid,
            no=no_counter,
            nama_dosen=nama_dosen,
            dtps=_b(row[2]),
            bk_ps_diakreditasi=_f(row[3]),
            bk_ps_lain_pt=_f(row[4]),
            bk_ps_luar_pt=_f(row[5]),
            bk_penelitian=_f(row[6]),
            bk_pkm=_f(row[7]),
            bk_tugas_tambahan=_f(row[8]),
        ))
        no_counter += 1
    db.flush()


# ---------------------------------------------------------------------------
# Section 4d – Publikasi DTPS Akademik (fixed rows)
# ---------------------------------------------------------------------------

def _import_sheet_4d(wb, db: Session, sid: uuid.UUID) -> None:
    if "4d" not in wb.sheetnames:
        return
    ws = wb["4d"]
    db.query(LkpsPublikasiIlmiah).filter(
        LkpsPublikasiIlmiah.submission_id == sid,
        LkpsPublikasiIlmiah.sumber == "dtps",
        LkpsPublikasiIlmiah.jenis_program == "akademik",
    ).delete()
    for row_num, kode_publikasi in _PUBLIKASI_ROW_AKADEMIK.items():
        row = ws[row_num]
        # C=ts2, D=ts1, E=ts
        if not _has_data(row[2].value, row[3].value, row[4].value):
            continue
        db.add(LkpsPublikasiIlmiah(
            submission_id=sid,
            sumber="dtps",
            jenis_program="akademik",
            kode_publikasi=kode_publikasi,
            ts2=_i(row[2].value),
            ts1=_i(row[3].value),
            ts=_i(row[4].value),
        ))
    db.flush()


# ---------------------------------------------------------------------------
# Section 4e – Publikasi DTPS Vokasi (fixed rows)
# ---------------------------------------------------------------------------

def _import_sheet_4e(wb, db: Session, sid: uuid.UUID) -> None:
    if "4e" not in wb.sheetnames:
        return
    ws = wb["4e"]
    db.query(LkpsPublikasiIlmiah).filter(
        LkpsPublikasiIlmiah.submission_id == sid,
        LkpsPublikasiIlmiah.sumber == "dtps",
        LkpsPublikasiIlmiah.jenis_program == "vokasi",
    ).delete()
    for row_num, kode_publikasi in _PUBLIKASI_ROW_VOKASI.items():
        row = ws[row_num]
        if not _has_data(row[2].value, row[3].value, row[4].value):
            continue
        db.add(LkpsPublikasiIlmiah(
            submission_id=sid,
            sumber="dtps",
            jenis_program="vokasi",
            kode_publikasi=kode_publikasi,
            ts2=_i(row[2].value),
            ts1=_i(row[3].value),
            ts=_i(row[4].value),
        ))
    db.flush()


# ---------------------------------------------------------------------------
# Sections 4f-1..4f-4 and 6e3-1..6e3-4 – Luaran Penelitian (dynamic rows)
# ---------------------------------------------------------------------------

def _import_luaran_sheets(wb, db: Session, sid: uuid.UUID) -> None:
    for sheet_name, (sumber, jenis_luaran) in _LUARAN_SHEET_META.items():
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        db.query(LkpsLuaranPenelitian).filter(
            LkpsLuaranPenelitian.submission_id == sid,
            LkpsLuaranPenelitian.sumber == sumber,
            LkpsLuaranPenelitian.jenis_luaran == jenis_luaran,
        ).delete()
        start_row = _LUARAN_START[sheet_name]
        for row in ws.iter_rows(min_row=start_row, values_only=True):
            # B=judul, C=tanggal, then varies
            judul = _s(row[1])
            if not judul:
                continue
            rec = LkpsLuaranPenelitian(
                submission_id=sid,
                sumber=sumber,
                jenis_luaran=jenis_luaran,
                judul=judul,
                tanggal=_d(row[2]),
            )
            if jenis_luaran == "paten":
                if sumber == "mahasiswa":
                    # D=status_mahasiswa, E=nomor_paten
                    rec.status_mahasiswa = _s(row[3])
                    rec.nomor_paten = _s(row[4])
                else:
                    # D=nomor_paten
                    rec.nomor_paten = _s(row[3])
            elif jenis_luaran == "hak_cipta":
                # D=nomor_hki
                rec.nomor_hki = _s(row[3])
            elif jenis_luaran == "teknologi":
                # D=status_tkt, E=nomor_sertifikat_tkt
                rec.status_tkt = _s(row[3])
                rec.nomor_sertifikat_tkt = _s(row[4])
            elif jenis_luaran == "buku":
                # D=nomor_isbn
                rec.nomor_isbn = _s(row[3])
            db.add(rec)
    db.flush()


# ---------------------------------------------------------------------------
# Section 4g – Produk/Jasa DTPS
# ---------------------------------------------------------------------------

def _import_sheet_4g(wb, db: Session, sid: uuid.UUID) -> None:
    if "4g" not in wb.sheetnames:
        return
    ws = wb["4g"]
    db.query(LkpsProdukJasa).filter(
        LkpsProdukJasa.submission_id == sid,
        LkpsProdukJasa.sumber == "dtps",
    ).delete()
    current_nama_pembuat = None
    for row in ws.iter_rows(min_row=6, values_only=True):
        # B=nama_pembuat, C=nama_produk_jasa, D=deskripsi, E=bukti
        if not _has_data(row[1], row[2], row[3], row[4]):
            continue
        nama_pembuat = _carry_forward(row[1], current_nama_pembuat)
        if not nama_pembuat:
            continue
        current_nama_pembuat = nama_pembuat
        db.add(LkpsProdukJasa(
            submission_id=sid,
            sumber="dtps",
            nama_pembuat=nama_pembuat,
            nama_produk_jasa=_s(row[2]),
            deskripsi=_s(row[3]),
            bukti=_s(row[4]),
        ))
    db.flush()


# ---------------------------------------------------------------------------
# Section 4h – Kinerja DTPS
# ---------------------------------------------------------------------------

def _import_sheet_4h(wb, db: Session, sid: uuid.UUID) -> None:
    if "4h" not in wb.sheetnames:
        return
    ws = wb["4h"]
    db.query(LkpsKinerjaDtps).filter(LkpsKinerjaDtps.submission_id == sid).delete()
    no_counter = 1
    for row in ws.iter_rows(min_row=16, values_only=True):
        # B=nama_dosen, C=ts2, D=ts1, E=ts, F=keterangan
        nama_dosen = _s(row[1])
        if not nama_dosen:
            continue
        db.add(LkpsKinerjaDtps(
            submission_id=sid,
            no=no_counter,
            nama_dosen=nama_dosen,
            ts2=_i(row[2]),
            ts1=_i(row[3]),
            ts=_i(row[4]),
            keterangan=_s(row[5]),
        ))
        no_counter += 1
    db.flush()


# ---------------------------------------------------------------------------
# Section 4i – Sitasi DTPS
# ---------------------------------------------------------------------------

def _import_sheet_4i(wb, db: Session, sid: uuid.UUID) -> None:
    if "4i" not in wb.sheetnames:
        return
    ws = wb["4i"]
    db.query(LkpsSitasiDtps).filter(LkpsSitasiDtps.submission_id == sid).delete()
    no_counter = 1
    current_nama_dosen = None
    for row in ws.iter_rows(min_row=6, values_only=True):
        # B=nama_dosen, C=judul_artikel, D=jumlah_sitasi
        if not _has_data(row[1], row[2], row[3]):
            continue
        nama_dosen = _carry_forward(row[1], current_nama_dosen)
        if not nama_dosen:
            continue
        current_nama_dosen = nama_dosen
        db.add(LkpsSitasiDtps(
            submission_id=sid,
            no=no_counter,
            nama_dosen=nama_dosen,
            judul_artikel=_s(row[2]),
            jumlah_sitasi=_i(row[3]),
        ))
        no_counter += 1
    db.flush()


# ---------------------------------------------------------------------------
# Section 4j – Rekognisi DTPS
# ---------------------------------------------------------------------------

def _import_sheet_4j(wb, db: Session, sid: uuid.UUID) -> None:
    if "4j" not in wb.sheetnames:
        return
    ws = wb["4j"]
    db.query(LkpsRekognisiDtps).filter(LkpsRekognisiDtps.submission_id == sid).delete()
    no_counter = 1
    current_nama_dosen = None
    current_bidang_keahlian = None
    current_tingkat = None
    for row in ws.iter_rows(min_row=12, values_only=True):
        # B=nama_dosen, C=bidang_keahlian, D=rekognisi, E=bukti_pendukung,
        # F=lokal(V), G=nasional(V), H=internasional(V), I=tahun
        # Export writes: write_tingkat_check(ws, row, rec.tingkat, "H", "G", "F")
        # So col_intr=H, col_nas=G, col_lokal=F
        if not _has_data(row[1], row[2], row[3], row[4], row[5], row[6], row[7], row[8]):
            continue
        nama_dosen = _carry_forward(row[1], current_nama_dosen)
        if not nama_dosen:
            continue
        current_nama_dosen = nama_dosen
        current_bidang_keahlian = _carry_forward(row[2], current_bidang_keahlian)
        # F=index5=lokal, G=index6=nasional, H=index7=internasional
        current_tingkat = _tingkat(row[7], row[6], row[5]) or current_tingkat  # intr=H, nas=G, lokal=F
        db.add(LkpsRekognisiDtps(
            submission_id=sid,
            no=no_counter,
            nama_dosen=nama_dosen,
            bidang_keahlian=current_bidang_keahlian,
            rekognisi=_s(row[3]),
            bukti_pendukung=_s(row[4]),
            tingkat=current_tingkat,
            tahun=_i(row[8]),  # col I
        ))
        no_counter += 1
    db.flush()


# ---------------------------------------------------------------------------
# Section 4k – Pembimbing Lapangan (PPI)
# ---------------------------------------------------------------------------

def _import_sheet_4k(wb, db: Session, sid: uuid.UUID) -> None:
    if "4k" not in wb.sheetnames:
        return
    ws = wb["4k"]
    db.query(LkpsPembimbingLapangan).filter(
        LkpsPembimbingLapangan.submission_id == sid
    ).delete()
    no_counter = 1
    for row in ws.iter_rows(min_row=6, values_only=True):
        # B=nama, C=industri, D=bidang_keinsinyuran, E=pengalaman_kerja_tahun,
        # F=pendidikan_tinggi, G=IPM(V), H=IPU(V), I=nomor_sip,
        # J=tanggal_berakhir_sip, K=jumlah_bimbingan_3tahun
        nama = _s(row[1])
        if not nama:
            continue
        # G=index6=IPM, H=index7=IPU
        if _b(row[6]):
            kategori_sip = "IPM"
        elif _b(row[7]):
            kategori_sip = "IPU"
        else:
            kategori_sip = None
        db.add(LkpsPembimbingLapangan(
            submission_id=sid,
            no=no_counter,
            nama=nama,
            industri=_s(row[2]),
            bidang_keinsinyuran=_s(row[3]),
            pengalaman_kerja_tahun=_i(row[4]),
            pendidikan_tinggi=_s(row[5]),
            kategori_sip=kategori_sip,
            nomor_sip=_s(row[8]),
            tanggal_berakhir_sip=_d(row[9]),
            jumlah_bimbingan_3tahun=_i(row[10]),
        ))
        no_counter += 1
    db.flush()


# ---------------------------------------------------------------------------
# Section 5a – Prasarana
# ---------------------------------------------------------------------------

def _import_sheet_5a(wb, db: Session, sid: uuid.UUID) -> None:
    if "5a" not in wb.sheetnames:
        return
    ws = wb["5a"]
    db.query(LkpsPrasarana).filter(LkpsPrasarana.submission_id == sid).delete()
    no_counter = 1
    current_nama_prasarana = None
    current_jumlah_prasarana = None
    for row in ws.iter_rows(min_row=10, values_only=True):
        # B=nama_prasarana, C=jumlah_prasarana, D=nama_sarana, E=jumlah_standar_minimal,
        # F=jumlah_dimiliki, G=kepemilikan_sendiri(V), H=kepemilikan_sewa(V),
        # I=kondisi_terawat(V), J=kondisi_tidak_terawat(V), K=logbook_ada(V), L=logbook_tidak_ada(V)
        if not _has_data(*row[1:12]):
            continue
        nama_prasarana = _carry_forward(row[1], current_nama_prasarana)
        if not nama_prasarana:
            continue
        current_nama_prasarana = nama_prasarana
        current_jumlah_prasarana = _i(row[2]) if row[2] is not None else current_jumlah_prasarana
        if _b(row[6]):
            kepemilikan = "sendiri"
        elif _b(row[7]):
            kepemilikan = "sewa"
        else:
            kepemilikan = None
        if _b(row[8]):
            kondisi = "terawat"
        elif _b(row[9]):
            kondisi = "tidak_terawat"
        else:
            kondisi = None
        if _b(row[10]):
            logbook = "ada"
        elif _b(row[11]):
            logbook = "tidak_ada"
        else:
            logbook = None
        db.add(LkpsPrasarana(
            submission_id=sid,
            no=no_counter,
            nama_prasarana=nama_prasarana,
            jumlah_prasarana=current_jumlah_prasarana,
            nama_sarana=_s(row[3]),
            jumlah_standar_minimal=_i(row[4]),
            jumlah_dimiliki=_i(row[5]),
            kepemilikan=kepemilikan,
            kondisi=kondisi,
            logbook=logbook,
        ))
        no_counter += 1
    db.flush()


# ---------------------------------------------------------------------------
# Section 5b – K3L Dokumen
# ---------------------------------------------------------------------------

def _import_sheet_5b(wb, db: Session, sid: uuid.UUID) -> None:
    if "5b" not in wb.sheetnames:
        return
    ws = wb["5b"]
    db.query(LkpsK3lDokumen).filter(LkpsK3lDokumen.submission_id == sid).delete()
    no_counter = 1
    for row in ws.iter_rows(min_row=17, values_only=True):
        # B=jenis_dokumen, C=jumlah, D=riwayat_pengesahan
        jenis_dokumen = _s(row[1])
        if not jenis_dokumen:
            continue
        db.add(LkpsK3lDokumen(
            submission_id=sid,
            no=no_counter,
            jenis_dokumen=jenis_dokumen,
            jumlah=_i(row[2]),
            riwayat_pengesahan=_s(row[3]),
        ))
        no_counter += 1
    db.flush()


# ---------------------------------------------------------------------------
# Section 5c – K3L Fasilitas
# ---------------------------------------------------------------------------

def _import_sheet_5c(wb, db: Session, sid: uuid.UUID) -> None:
    if "5c" not in wb.sheetnames:
        return
    ws = wb["5c"]
    db.query(LkpsK3lFasilitas).filter(LkpsK3lFasilitas.submission_id == sid).delete()
    no_counter = 1
    for row in ws.iter_rows(min_row=17, values_only=True):
        # B=nama_sarana, C=fungsi, D=jumlah_unit, E=kondisi_terawat(V), F=kondisi_tidak_terawat(V)
        nama_sarana = _s(row[1])
        if not nama_sarana:
            continue
        if _b(row[4]):
            kondisi = "terawat"
        elif _b(row[5]):
            kondisi = "tidak_terawat"
        else:
            kondisi = None
        db.add(LkpsK3lFasilitas(
            submission_id=sid,
            no=no_counter,
            nama_sarana=nama_sarana,
            fungsi=_s(row[2]),
            jumlah_unit=_i(row[3]),
            kondisi=kondisi,
        ))
        no_counter += 1
    db.flush()


# ---------------------------------------------------------------------------
# Section 6a – Mahasiswa Aktif
# ---------------------------------------------------------------------------

def _import_sheet_6a(wb, db: Session, sid: uuid.UUID) -> None:
    if "6a" not in wb.sheetnames:
        return
    ws = wb["6a"]
    db.query(LkpsMahasiswaAktif).filter(LkpsMahasiswaAktif.submission_id == sid).delete()
    no_counter = 1
    for row in ws.iter_rows(min_row=7, values_only=True):
        # B=program_studi_nama, C=prodi_diakreditasi(√/V), D=aktif_ts2, E=aktif_ts1, F=aktif_ts,
        # G=asing_fulltime_ts2, H=asing_fulltime_ts1, I=asing_fulltime_ts,
        # J=asing_parttime_ts2, K=asing_parttime_ts1, L=asing_parttime_ts
        prodi_nama = _s(row[1])
        if not prodi_nama:
            continue
        # prodi_diakreditasi: the export writes "√", but we accept "V" or "√"
        prodi_diakreditasi_val = row[2]
        prodi_diakreditasi = (
            str(prodi_diakreditasi_val).strip() in ("√", "V", "v")
            if prodi_diakreditasi_val is not None else False
        )
        db.add(LkpsMahasiswaAktif(
            submission_id=sid,
            no=no_counter,
            program_studi_nama=prodi_nama,
            prodi_diakreditasi=prodi_diakreditasi,
            aktif_ts2=_i(row[3]),
            aktif_ts1=_i(row[4]),
            aktif_ts=_i(row[5]),
            asing_fulltime_ts2=_i(row[6]),
            asing_fulltime_ts1=_i(row[7]),
            asing_fulltime_ts=_i(row[8]),
            asing_parttime_ts2=_i(row[9]),
            asing_parttime_ts1=_i(row[10]),
            asing_parttime_ts=_i(row[11]),
        ))
        no_counter += 1
    db.flush()


# ---------------------------------------------------------------------------
# Section 6b – IPK Lulusan (fixed rows)
# ---------------------------------------------------------------------------

def _import_sheet_6b(wb, db: Session, sid: uuid.UUID) -> None:
    if "6b" not in wb.sheetnames:
        return
    ws = wb["6b"]
    db.query(LkpsIpkLulusan).filter(LkpsIpkLulusan.submission_id == sid).delete()
    for row_num, periode in _IPK_ROW_INV.items():
        row = ws[row_num]
        # C=jumlah_lulusan, D=ipk_min, E=ipk_rata, F=ipk_maks
        if not _has_data(row[2].value, row[3].value, row[4].value, row[5].value):
            continue
        db.add(LkpsIpkLulusan(
            submission_id=sid,
            periode=periode,
            jumlah_lulusan=_i(row[2].value),
            ipk_min=_f(row[3].value),
            ipk_rata=_f(row[4].value),
            ipk_maks=_f(row[5].value),
        ))
    db.flush()


# ---------------------------------------------------------------------------
# Section 6c1 / 6c2 – Prestasi Mahasiswa
# ---------------------------------------------------------------------------

_PRESTASI_META = [
    ("akademik",     "6c1", 10),
    ("non_akademik", "6c2", 11),
]


def _import_sheet_6c(wb, db: Session, sid: uuid.UUID) -> None:
    for jenis, sheet_name, start_row in _PRESTASI_META:
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        db.query(LkpsPrestasiMahasiswa).filter(
            LkpsPrestasiMahasiswa.submission_id == sid,
            LkpsPrestasiMahasiswa.jenis == jenis,
        ).delete()
        no_counter = 1
        for row in ws.iter_rows(min_row=start_row, values_only=True):
            # B=nama_kegiatan, C=waktu_perolehan, D=lokal(V), E=nasional(V), F=internasional(V), G=prestasi_dicapai
            # Export writes: write_tingkat_check(ws, row, rec.tingkat, "F", "E", "D")
            # So col_intr=F, col_nas=E, col_lokal=D → index5=F, index4=E, index3=D
            nama_kegiatan = _s(row[1])
            if not nama_kegiatan:
                continue
            tingkat = _tingkat(row[5], row[4], row[3])  # F=intr, E=nas, D=lokal
            db.add(LkpsPrestasiMahasiswa(
                submission_id=sid,
                jenis=jenis,
                no=no_counter,
                nama_kegiatan=nama_kegiatan,
                waktu_perolehan=_d(row[2]),
                tingkat=tingkat,
                prestasi_dicapai=_s(row[6]),
            ))
            no_counter += 1
    db.flush()


# ---------------------------------------------------------------------------
# Section 6d – Masa Studi Lulusan
# ---------------------------------------------------------------------------

def _import_sheet_6d(wb, db: Session, sid: uuid.UUID) -> None:
    if "6d" not in wb.sheetnames:
        return
    ws = wb["6d"]
    db.query(LkpsMasaStudi).filter(LkpsMasaStudi.submission_id == sid).delete()
    for jenis_program, (first_row, tahun_labels) in _MASA_STUDI_SECTIONS.items():
        for offset, tahun_masuk in enumerate(tahun_labels):
            row_num = first_row + offset
            row = ws[row_num]
            # B=jumlah_masuk, C=lulus_tepat_waktu, D=lulus_terlambat, E=tidak_lulus
            if not _has_data(row[1].value, row[2].value, row[3].value, row[4].value):
                continue
            db.add(LkpsMasaStudi(
                submission_id=sid,
                jenis_program=jenis_program,
                tahun_masuk=tahun_masuk,
                jumlah_masuk=_i(row[1].value),
                jumlah_lulus_tepat_waktu=_i(row[2].value),
                jumlah_lulus_terlambat=_i(row[3].value),
                jumlah_tidak_lulus=_i(row[4].value),
            ))
    db.flush()


# ---------------------------------------------------------------------------
# Section 6e1 – Publikasi Mahasiswa Akademik (fixed rows)
# ---------------------------------------------------------------------------

def _import_sheet_6e1(wb, db: Session, sid: uuid.UUID) -> None:
    if "6e1" not in wb.sheetnames:
        return
    ws = wb["6e1"]
    db.query(LkpsPublikasiIlmiah).filter(
        LkpsPublikasiIlmiah.submission_id == sid,
        LkpsPublikasiIlmiah.sumber == "mahasiswa",
        LkpsPublikasiIlmiah.jenis_program == "akademik",
    ).delete()
    for row_num, kode_publikasi in _PUBLIKASI_ROW_AKADEMIK.items():
        row = ws[row_num]
        if not _has_data(row[2].value, row[3].value, row[4].value):
            continue
        db.add(LkpsPublikasiIlmiah(
            submission_id=sid,
            sumber="mahasiswa",
            jenis_program="akademik",
            kode_publikasi=kode_publikasi,
            ts2=_i(row[2].value),
            ts1=_i(row[3].value),
            ts=_i(row[4].value),
        ))
    db.flush()


# ---------------------------------------------------------------------------
# Section 6e2 – Publikasi Mahasiswa Vokasi (fixed rows)
# ---------------------------------------------------------------------------

def _import_sheet_6e2(wb, db: Session, sid: uuid.UUID) -> None:
    if "6e2" not in wb.sheetnames:
        return
    ws = wb["6e2"]
    db.query(LkpsPublikasiIlmiah).filter(
        LkpsPublikasiIlmiah.submission_id == sid,
        LkpsPublikasiIlmiah.sumber == "mahasiswa",
        LkpsPublikasiIlmiah.jenis_program == "vokasi",
    ).delete()
    for row_num, kode_publikasi in _PUBLIKASI_ROW_VOKASI.items():
        row = ws[row_num]
        if not _has_data(row[2].value, row[3].value, row[4].value):
            continue
        db.add(LkpsPublikasiIlmiah(
            submission_id=sid,
            sumber="mahasiswa",
            jenis_program="vokasi",
            kode_publikasi=kode_publikasi,
            ts2=_i(row[2].value),
            ts1=_i(row[3].value),
            ts=_i(row[4].value),
        ))
    db.flush()


# ---------------------------------------------------------------------------
# Section 6e4 – Produk/Jasa Mahasiswa
# ---------------------------------------------------------------------------

def _import_sheet_6e4(wb, db: Session, sid: uuid.UUID) -> None:
    if "6e4" not in wb.sheetnames:
        return
    ws = wb["6e4"]
    db.query(LkpsProdukJasa).filter(
        LkpsProdukJasa.submission_id == sid,
        LkpsProdukJasa.sumber == "mahasiswa",
    ).delete()
    current_nama_pembuat = None
    for row in ws.iter_rows(min_row=6, values_only=True):
        # B=nama_pembuat, C=nama_produk_jasa, D=deskripsi, E=bukti
        if not _has_data(row[1], row[2], row[3], row[4]):
            continue
        nama_pembuat = _carry_forward(row[1], current_nama_pembuat)
        if not nama_pembuat:
            continue
        current_nama_pembuat = nama_pembuat
        db.add(LkpsProdukJasa(
            submission_id=sid,
            sumber="mahasiswa",
            nama_pembuat=nama_pembuat,
            nama_produk_jasa=_s(row[2]),
            deskripsi=_s(row[3]),
            bukti=_s(row[4]),
        ))
    db.flush()


# ---------------------------------------------------------------------------
# Section 6f1 – Waktu Tunggu (fixed-section rows)
# ---------------------------------------------------------------------------

def _import_sheet_6f1(wb, db: Session, sid: uuid.UUID) -> None:
    if "6f1" not in wb.sheetnames:
        return
    ws = wb["6f1"]
    db.query(LkpsWaktuTunggu).filter(LkpsWaktuTunggu.submission_id == sid).delete()
    for jenis_program, section_start in _WT_SECTION_START.items():
        for offset, tahun_lulus in _WT_TAHUN_OFFSET_INV.items():
            row_num = section_start + offset
            row = ws[row_num]
            # B=jumlah_lulusan, C=jumlah_terlacak, D=jumlah_dipesan_sebelum_lulus,
            # E=wt_lt_3bulan, F=wt_3_6bulan, G=wt_gt_6bulan
            if not _has_data(row[1].value, row[2].value, row[3].value, row[4].value, row[5].value, row[6].value):
                continue
            db.add(LkpsWaktuTunggu(
                submission_id=sid,
                jenis_program=jenis_program,
                tahun_lulus=tahun_lulus,
                jumlah_lulusan=_i(row[1].value),
                jumlah_terlacak=_i(row[2].value),
                jumlah_dipesan_sebelum_lulus=_i(row[3].value),
                wt_lt_3bulan=_i(row[4].value),
                wt_3_6bulan=_i(row[5].value),
                wt_gt_6bulan=_i(row[6].value),
            ))
    db.flush()


# ---------------------------------------------------------------------------
# Section 6f2 – Kesesuaian Kerja (fixed rows)
# ---------------------------------------------------------------------------

def _import_sheet_6f2(wb, db: Session, sid: uuid.UUID) -> None:
    if "6f2" not in wb.sheetnames:
        return
    ws = wb["6f2"]
    db.query(LkpsKesesuaianKerja).filter(
        LkpsKesesuaianKerja.submission_id == sid
    ).delete()
    for row_num, tahun_lulus in _LULUSAN_ROW_INV.items():
        row = ws[row_num]
        # B=jumlah_lulusan, C=jumlah_terlacak, D=kesesuaian_rendah,
        # E=kesesuaian_sedang, F=kesesuaian_tinggi
        if not _has_data(row[1].value, row[2].value, row[3].value, row[4].value, row[5].value):
            continue
        db.add(LkpsKesesuaianKerja(
            submission_id=sid,
            tahun_lulus=tahun_lulus,
            jumlah_lulusan=_i(row[1].value),
            jumlah_terlacak=_i(row[2].value),
            kesesuaian_rendah=_i(row[3].value),
            kesesuaian_sedang=_i(row[4].value),
            kesesuaian_tinggi=_i(row[5].value),
        ))
    db.flush()


# ---------------------------------------------------------------------------
# Section 6g1 – Tempat Kerja (fixed rows)
# ---------------------------------------------------------------------------

def _import_sheet_6g1(wb, db: Session, sid: uuid.UUID) -> None:
    if "6g1" not in wb.sheetnames:
        return
    ws = wb["6g1"]
    db.query(LkpsTempatKerja).filter(LkpsTempatKerja.submission_id == sid).delete()
    for row_num, tahun_lulus in _LULUSAN_ROW_INV.items():
        row = ws[row_num]
        # B=jumlah_lulusan, C=jumlah_pengguna_tanggapan, D=jumlah_terlacak,
        # E=bekerja_lokal, F=bekerja_nasional, G=bekerja_multinasional
        if not _has_data(row[1].value, row[2].value, row[3].value, row[4].value, row[5].value, row[6].value):
            continue
        db.add(LkpsTempatKerja(
            submission_id=sid,
            tahun_lulus=tahun_lulus,
            jumlah_lulusan=_i(row[1].value),
            jumlah_pengguna_tanggapan=_i(row[2].value),
            jumlah_terlacak=_i(row[3].value),
            bekerja_lokal=_i(row[4].value),
            bekerja_nasional=_i(row[5].value),
            bekerja_multinasional=_i(row[6].value),
        ))
    db.flush()


# ---------------------------------------------------------------------------
# Section 6g2 – Kepuasan Pengguna (fixed rows)
# ---------------------------------------------------------------------------

def _import_sheet_6g2(wb, db: Session, sid: uuid.UUID) -> None:
    if "6g2" not in wb.sheetnames:
        return
    ws = wb["6g2"]
    db.query(LkpsKepuasanPengguna).filter(
        LkpsKepuasanPengguna.submission_id == sid
    ).delete()
    for row_num, jenis_kemampuan in _KEPUASAN_ROW_INV.items():
        row = ws[row_num]
        # C=sangat_baik, D=baik, E=cukup, F=kurang, G=rencana_tindak_lanjut
        if not _has_data(row[2].value, row[3].value, row[4].value, row[5].value, row[6].value):
            continue
        db.add(LkpsKepuasanPengguna(
            submission_id=sid,
            no=row_num - 6,  # row 7 → no=1, row 13 → no=7
            jenis_kemampuan=jenis_kemampuan,
            sangat_baik=_f(row[2].value),
            baik=_f(row[3].value),
            cukup=_f(row[4].value),
            kurang=_f(row[5].value),
            rencana_tindak_lanjut=_s(row[6].value),
        ))
    db.flush()


# ---------------------------------------------------------------------------
# Section 6h1 – Penelitian Mahasiswa
# ---------------------------------------------------------------------------

def _import_sheet_6h1(wb, db: Session, sid: uuid.UUID) -> None:
    if "6h1" not in wb.sheetnames:
        return
    ws = wb["6h1"]
    db.query(LkpsPenelitianMahasiswa).filter(
        LkpsPenelitianMahasiswa.submission_id == sid,
        LkpsPenelitianMahasiswa.jenis == "penelitian",
    ).delete()
    no_counter = 1
    current_nama_dosen = None
    current_tema_penelitian = None
    for row in ws.iter_rows(min_row=11, values_only=True):
        # B=nama_dosen, C=tema_penelitian, D=nama_mahasiswa, E=judul_kegiatan, F=tahun
        if not _has_data(*row[1:6]):
            continue
        nama_dosen = _carry_forward(row[1], current_nama_dosen)
        if not nama_dosen:
            continue
        current_nama_dosen = nama_dosen
        current_tema_penelitian = _carry_forward(row[2], current_tema_penelitian)
        db.add(LkpsPenelitianMahasiswa(
            submission_id=sid,
            jenis="penelitian",
            no=no_counter,
            nama_dosen=nama_dosen,
            tema_penelitian=current_tema_penelitian,
            nama_mahasiswa=_s(row[3]),
            judul_kegiatan=_s(row[4]),
            tahun=_i(row[5]),
        ))
        no_counter += 1
    db.flush()


# ---------------------------------------------------------------------------
# Section 6h2 – Tesis/Disertasi Mahasiswa
# ---------------------------------------------------------------------------

def _import_sheet_6h2(wb, db: Session, sid: uuid.UUID) -> None:
    if "6h2" not in wb.sheetnames:
        return
    ws = wb["6h2"]
    db.query(LkpsPenelitianMahasiswa).filter(
        LkpsPenelitianMahasiswa.submission_id == sid,
        LkpsPenelitianMahasiswa.jenis == "tesis_disertasi",
    ).delete()
    no_counter = 1
    current_nama_dosen = None
    current_tema_penelitian = None
    for row in ws.iter_rows(min_row=6, values_only=True):
        if not _has_data(*row[1:6]):
            continue
        nama_dosen = _carry_forward(row[1], current_nama_dosen)
        if not nama_dosen:
            continue
        current_nama_dosen = nama_dosen
        current_tema_penelitian = _carry_forward(row[2], current_tema_penelitian)
        db.add(LkpsPenelitianMahasiswa(
            submission_id=sid,
            jenis="tesis_disertasi",
            no=no_counter,
            nama_dosen=nama_dosen,
            tema_penelitian=current_tema_penelitian,
            nama_mahasiswa=_s(row[3]),
            judul_kegiatan=_s(row[4]),
            tahun=_i(row[5]),
        ))
        no_counter += 1
    db.flush()


# ---------------------------------------------------------------------------
# Section 6i – PkM Mahasiswa
# ---------------------------------------------------------------------------

def _import_sheet_6i(wb, db: Session, sid: uuid.UUID) -> None:
    if "6i" not in wb.sheetnames:
        return
    ws = wb["6i"]
    db.query(LkpsPenelitianMahasiswa).filter(
        LkpsPenelitianMahasiswa.submission_id == sid,
        LkpsPenelitianMahasiswa.jenis == "pkm",
    ).delete()
    no_counter = 1
    current_nama_dosen = None
    current_tema_penelitian = None
    for row in ws.iter_rows(min_row=6, values_only=True):
        if not _has_data(*row[1:6]):
            continue
        nama_dosen = _carry_forward(row[1], current_nama_dosen)
        if not nama_dosen:
            continue
        current_nama_dosen = nama_dosen
        current_tema_penelitian = _carry_forward(row[2], current_tema_penelitian)
        db.add(LkpsPenelitianMahasiswa(
            submission_id=sid,
            jenis="pkm",
            no=no_counter,
            nama_dosen=nama_dosen,
            tema_penelitian=current_tema_penelitian,
            nama_mahasiswa=_s(row[3]),
            judul_kegiatan=_s(row[4]),
            tahun=_i(row[5]),
        ))
        no_counter += 1
    db.flush()


# ---------------------------------------------------------------------------
# Section 7a – SPMI Dokumen (fixed rows)
# ---------------------------------------------------------------------------

def _import_sheet_7a(wb, db: Session, sid: uuid.UUID) -> None:
    if "7a" not in wb.sheetnames:
        return
    ws = wb["7a"]
    db.query(LkpsSpmiDokumen).filter(LkpsSpmiDokumen.submission_id == sid).delete()
    for row_num, jenis_dokumen in _SPMI_DOK_ROW_INV.items():
        row = ws[row_num]
        # B=pre-labeled (jenis_dokumen), C=no_dokumen, D=tanggal_dokumen
        if not _has_data(row[2].value, row[3].value):
            continue
        db.add(LkpsSpmiDokumen(
            submission_id=sid,
            no=row_num - 4,  # row 5 → no=1
            jenis_dokumen=jenis_dokumen,
            no_dokumen=_s(row[2].value),
            tanggal_dokumen=_d(row[3].value),
        ))
    db.flush()


# ---------------------------------------------------------------------------
# Section 7b – SPMI Pelaksanaan (fixed rows)
# ---------------------------------------------------------------------------

def _import_sheet_7b(wb, db: Session, sid: uuid.UUID) -> None:
    if "7b" not in wb.sheetnames:
        return
    ws = wb["7b"]
    db.query(LkpsSpmiPelaksanaan).filter(
        LkpsSpmiPelaksanaan.submission_id == sid
    ).delete()
    for row_num, jenis_pelaksanaan in _PPEPP_ROW_INV.items():
        row = ws[row_num]
        # C=link_dokumen, D=link_laporan_audit, E=link_laporan_rtm, F=link_dokumen_peningkatan
        if not _has_data(row[2].value, row[3].value, row[4].value, row[5].value):
            continue
        db.add(LkpsSpmiPelaksanaan(
            submission_id=sid,
            no=row_num - 4,  # row 5 → no=1
            jenis_pelaksanaan=jenis_pelaksanaan,
            link_dokumen=_s(row[2].value),
            link_laporan_audit=_s(row[3].value),
            link_laporan_rtm=_s(row[4].value),
            link_dokumen_peningkatan=_s(row[5].value),
        ))
    db.flush()
